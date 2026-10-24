import io
from typing import Dict, Any, List, Union
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


HEADER_FILL  = PatternFill("solid", fgColor="1E3A8A")
SUBHDR_FILL  = PatternFill("solid", fgColor="2563EB")
EVEN_FILL    = PatternFill("solid", fgColor="F8FAFC")
ODD_FILL     = PatternFill("solid", fgColor="FFFFFF")
TOTAL_FILL   = PatternFill("solid", fgColor="DBEAFE")

HEADER_FONT  = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
SUBHDR_FONT  = Font(name="Segoe UI", size=9,  bold=True, color="FFFFFF")
BODY_FONT    = Font(name="Segoe UI", size=9,  color="334155")
TOTAL_FONT   = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
MONO_FONT    = Font(name="Courier New", size=9, color="334155")

PAID_FILL    = PatternFill("solid", fgColor="DCFCE7")   # Light green
PARTIAL_FILL = PatternFill("solid", fgColor="FEF9C3")   # Light yellow
UNPAID_FILL  = PatternFill("solid", fgColor="FEE2E2")   # Light red
PAID_FONT    = Font(name="Segoe UI", size=8, bold=True, color="15803D")
PARTIAL_FONT = Font(name="Segoe UI", size=8, bold=True, color="92400E")
UNPAID_FONT  = Font(name="Segoe UI", size=8, bold=True, color="DC2626")

NUM_FMT = "#,##0.00"
thin = Side(border_style="thin", color="CBD5E1")
CELL_BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

CENTER = Alignment(horizontal="center", vertical="center")
RIGHT  = Alignment(horizontal="right",  vertical="center")
LEFT   = Alignment(horizontal="left",   vertical="center")


def _apply_kpi_cards(ws, summaries: List[tuple], start_row: int = 3, num_cols: int = 8):
    """Draw clean KPI cards at the top of an Excel sheet."""
    # We place 4 KPI cards across columns A through H (2 columns per card)
    col_starts = [1, 3, 5, 7]
    for i, (label, val_str, bg_color) in enumerate(summaries[:4]):
        c_start = col_starts[i]
        c_end = c_start + 1
        c_start_let = get_column_letter(c_start)
        c_end_let = get_column_letter(c_end)

        # Label row
        ws.merge_cells(f"{c_start_let}{start_row}:{c_end_let}{start_row}")
        lc = ws.cell(row=start_row, column=c_start, value=label)
        lc.font = Font(name="Segoe UI", size=8, bold=True, color="FFFFFF")
        lc.fill = PatternFill("solid", fgColor=bg_color)
        lc.alignment = CENTER; lc.border = CELL_BORDER

        # Value row
        ws.merge_cells(f"{c_start_let}{start_row+1}:{c_end_let}{start_row+1}")
        vc = ws.cell(row=start_row+1, column=c_start, value=val_str)
        vc.font = Font(name="Segoe UI", size=12, bold=True, color=bg_color)
        vc.fill = PatternFill("solid", fgColor="F8FAFC")
        vc.alignment = CENTER; vc.border = CELL_BORDER


def build_daily_by_customer(report_data: Union[Dict[str, Any], List[Dict]], report_date: str) -> io.BytesIO:
    """
    Builds a comprehensive Daily Sales & Customer Report with two sheets:
      1. Customer Summary: Every customer with Total Sale, Received, Pending, and Status
      2. Detailed Bills & Items: Itemized bill breakdown per customer with subtotals
    """
    wb = Workbook()

    # Normalize data format
    if isinstance(report_data, dict):
        customers = report_data.get("customers", [])
        detailed_rows = report_data.get("detailed_rows", [])
        totals = report_data.get("totals", {})
    else:
        # Backward compatibility if a raw list of detailed rows was provided
        detailed_rows = report_data
        customers_map = {}
        for r in detailed_rows:
            cname = r.get("customer_name") or "Unknown"
            if cname not in customers_map:
                customers_map[cname] = {
                    "customer_name": cname,
                    "phone": r.get("customer_phone") or "",
                    "bills_count": 0,
                    "total_sale": 0.0,
                    "amount_paid": 0.0,
                    "balance_due": 0.0,
                    "_seen_bills": set(),
                }
            cm = customers_map[cname]
            b_no = r.get("bill_no")
            if b_no and b_no not in cm["_seen_bills"]:
                cm["_seen_bills"].add(b_no)
                cm["bills_count"] += 1
                cm["total_sale"] += float(r.get("total_bill_amount") or r.get("amount") or 0)
                cm["amount_paid"] += float(r.get("amount_paid") or 0)
                cm["balance_due"] += float(r.get("balance_due") or 0)

        customers = []
        for cm in customers_map.values():
            paid = cm["amount_paid"]
            due = cm["balance_due"]
            tot = cm["total_sale"]
            status = "paid" if tot > 0 and due <= 0 else "partial" if paid > 0 else "unpaid"
            customers.append({
                "customer_name": cm["customer_name"],
                "phone": cm["phone"],
                "bills_count": cm["bills_count"],
                "total_sale": tot,
                "amount_paid": paid,
                "balance_due": due,
                "payment_status": status,
            })
        totals = {
            "total_sales": sum(c["total_sale"] for c in customers),
            "total_received": sum(c["amount_paid"] for c in customers),
            "total_pending": sum(c["balance_due"] for c in customers),
            "total_bills": sum(c["bills_count"] for c in customers),
        }

    grand_sale = float(totals.get("total_sales") or sum(float(c.get("total_sale") or 0) for c in customers))
    grand_received = float(totals.get("total_received") or sum(float(c.get("amount_paid") or 0) for c in customers))
    grand_pending = float(totals.get("total_pending") or sum(float(c.get("balance_due") or 0) for c in customers))
    collection_rate = (grand_received / grand_sale * 100) if grand_sale > 0 else 0.0

    # ══════════════════════════════════════════════════════════════════════════
    # SHEET 1: CUSTOMER SUMMARY
    # ══════════════════════════════════════════════════════════════════════════
    ws1 = wb.active
    ws1.title = "Customer Summary"

    # Banner
    ws1.merge_cells("A1:I1")
    t1 = ws1["A1"]
    t1.value = f"DAILY SALES & PAYMENT REPORT — CUSTOMER SUMMARY   |   {report_date}"
    t1.font = HEADER_FONT; t1.fill = HEADER_FILL; t1.alignment = CENTER
    ws1.row_dimensions[1].height = 32

    # KPI Cards (Rows 3 & 4)
    kpis = [
        ("TOTAL SALES", f"₹{grand_sale:,.2f}", "16A34A"),
        ("RECEIVED AMOUNT", f"₹{grand_received:,.2f}", "2563EB"),
        ("PENDING BALANCE", f"₹{grand_pending:,.2f}", "DC2626"),
        ("COLLECTION RATE", f"{collection_rate:.1f}%", "7C3AED"),
    ]
    _apply_kpi_cards(ws1, kpis, start_row=3, num_cols=8)
    ws1.row_dimensions[3].height = 18
    ws1.row_dimensions[4].height = 24

    # Table Header (Row 6)
    headers1 = [
        "#", "Customer Name", "Phone", "Bills Count",
        "Total Sale (₹)", "Received (₹)", "Pending (₹)", "Profit / Loss (₹)", "Payment Status"
    ]
    for ci, h in enumerate(headers1, 1):
        cell = ws1.cell(row=6, column=ci, value=h)
        cell.font = SUBHDR_FONT; cell.fill = SUBHDR_FILL
        cell.alignment = CENTER; cell.border = CELL_BORDER
    ws1.row_dimensions[6].height = 22

    # Data Rows
    current_row = 7
    for idx, c in enumerate(customers, 1):
        fill = EVEN_FILL if current_row % 2 == 0 else ODD_FILL
        c_sale = float(c.get("total_sale") or 0)
        c_paid = float(c.get("amount_paid") or 0)
        c_due = float(c.get("balance_due") or 0)
        pl = c.get("profit_loss")
        status = (c.get("payment_status") or "unpaid").lower()

        ws1.cell(row=current_row, column=1, value=idx).alignment = CENTER
        ws1.cell(row=current_row, column=2, value=c.get("customer_name") or "—").alignment = LEFT
        ws1.cell(row=current_row, column=3, value=c.get("phone") or "—").alignment = CENTER
        ws1.cell(row=current_row, column=4, value=c.get("bills_count") or 1).alignment = CENTER

        sc = ws1.cell(row=current_row, column=5, value=c_sale)
        sc.number_format = NUM_FMT; sc.alignment = RIGHT; sc.font = Font(name="Segoe UI", size=9, bold=True, color="166534")

        pc = ws1.cell(row=current_row, column=6, value=c_paid)
        pc.number_format = NUM_FMT; pc.alignment = RIGHT; pc.font = Font(name="Segoe UI", size=9, bold=True, color="1E40AF")

        dc = ws1.cell(row=current_row, column=7, value=c_due)
        dc.number_format = NUM_FMT; dc.alignment = RIGHT; dc.font = Font(name="Segoe UI", size=9, bold=True, color="991B1B")

        plc = ws1.cell(row=current_row, column=8, value=pl if pl is not None else "—")
        if pl is not None:
            plc.number_format = NUM_FMT; plc.alignment = RIGHT
            plc.font = Font(name="Segoe UI", size=9, bold=True, color="166534" if pl >= 0 else "991B1B")
        else:
            plc.alignment = CENTER; plc.font = BODY_FONT

        st_cell = ws1.cell(row=current_row, column=9, value=status.upper())
        st_cell.alignment = CENTER
        if status == "paid":
            st_cell.fill = PAID_FILL; st_cell.font = PAID_FONT
        elif status == "partial":
            st_cell.fill = PARTIAL_FILL; st_cell.font = PARTIAL_FONT
        else:
            st_cell.fill = UNPAID_FILL; st_cell.font = UNPAID_FONT

        for ci in range(1, 10):
            cell = ws1.cell(row=current_row, column=ci)
            cell.border = CELL_BORDER
            if ci != 9:
                cell.fill = fill
                if ci < 5:
                    cell.font = BODY_FONT

        ws1.row_dimensions[current_row].height = 20
        current_row += 1

    # Grand Total Row
    ws1.merge_cells(f"A{current_row}:D{current_row}")
    gt_lbl = ws1.cell(row=current_row, column=1, value="GRAND TOTAL")
    gt_lbl.font = TOTAL_FONT; gt_lbl.fill = TOTAL_FILL
    gt_lbl.alignment = RIGHT; gt_lbl.border = CELL_BORDER

    for ci, val in enumerate([grand_sale, grand_received, grand_pending], start=5):
        cell = ws1.cell(row=current_row, column=ci, value=val)
        cell.font = TOTAL_FONT; cell.fill = TOTAL_FILL
        cell.number_format = NUM_FMT; cell.alignment = RIGHT
        cell.border = CELL_BORDER

    has_c_pl = any(c.get("profit_loss") is not None for c in customers)
    grand_c_pl = sum(c["profit_loss"] for c in customers if c.get("profit_loss") is not None)
    pl_total = ws1.cell(row=current_row, column=8, value=grand_c_pl if has_c_pl else "—")
    pl_total.font = TOTAL_FONT; pl_total.fill = TOTAL_FILL; pl_total.border = CELL_BORDER
    if has_c_pl:
        pl_total.number_format = NUM_FMT; pl_total.alignment = RIGHT
    else:
        pl_total.alignment = CENTER

    st_total = ws1.cell(row=current_row, column=9, value=f"{collection_rate:.1f}% Recv")
    st_total.font = TOTAL_FONT; st_total.fill = TOTAL_FILL
    st_total.alignment = CENTER; st_total.border = CELL_BORDER
    ws1.row_dimensions[current_row].height = 24

    # Column Widths for Sheet 1
    w1 = [6, 28, 16, 14, 18, 18, 18, 18, 16]
    for ci, w in enumerate(w1, 1):
        ws1.column_dimensions[get_column_letter(ci)].width = w

    # ══════════════════════════════════════════════════════════════════════════
    # SHEET 2: DETAILED BILLS & ITEMS
    # ══════════════════════════════════════════════════════════════════════════
    ws2 = wb.create_sheet(title="Detailed Bills & Items")

    headers2 = [
        "Customer", "Phone", "Bill No", "Date",
        "Product", "Qty", "Unit", "Rate (₹)", "Item Total (₹)",
        "Bill Total (₹)", "Received (₹)", "Pending (₹)", "Payment Method", "Status"
    ]
    NCOLS2 = len(headers2)

    # Banner
    ws2.merge_cells(f"A1:{get_column_letter(NCOLS2)}1")
    t2 = ws2["A1"]
    t2.value = f"DETAILED BILLS & ITEM BREAKDOWN   |   {report_date}"
    t2.font = HEADER_FONT; t2.fill = HEADER_FILL; t2.alignment = CENTER
    ws2.row_dimensions[1].height = 30

    for ci, h in enumerate(headers2, 1):
        cell = ws2.cell(row=2, column=ci, value=h)
        cell.font = SUBHDR_FONT; cell.fill = SUBHDR_FILL
        cell.alignment = CENTER; cell.border = CELL_BORDER
    ws2.row_dimensions[2].height = 20

    r2 = 3
    prev_cust = None
    sub_sale = sub_paid = sub_due = 0.0
    seen_bills_sub = set()

    for row in detailed_rows:
        cname = row.get("customer_name") or "Unknown"
        b_no = row.get("bill_no") or ""

        # Customer Subtotal separator
        if prev_cust is not None and cname != prev_cust:
            ws2.merge_cells(f"A{r2}:I{r2}")
            sub_lbl = ws2.cell(row=r2, column=1, value=f"  ↳ Total for {prev_cust}")
            sub_lbl.font = Font(name="Segoe UI", size=9, bold=True, color="1E40AF")
            sub_lbl.fill = PatternFill("solid", fgColor="EFF6FF")
            sub_lbl.alignment = LEFT; sub_lbl.border = CELL_BORDER

            for ci, val in enumerate([sub_sale, sub_paid, sub_due], start=10):
                cell = ws2.cell(row=r2, column=ci, value=val)
                cell.font = Font(name="Segoe UI", size=9, bold=True, color="1E40AF")
                cell.fill = PatternFill("solid", fgColor="EFF6FF")
                cell.number_format = NUM_FMT; cell.alignment = RIGHT; cell.border = CELL_BORDER

            for ci in [13, 14]:
                c_end = ws2.cell(row=r2, column=ci, value="")
                c_end.fill = PatternFill("solid", fgColor="EFF6FF"); c_end.border = CELL_BORDER

            r2 += 1
            sub_sale = sub_paid = sub_due = 0.0
            seen_bills_sub.clear()

        # Track bill financial totals once per bill
        t_bill = row.get("total_bill_amount")
        a_paid = row.get("amount_paid")
        b_due = row.get("balance_due")
        is_first_item = row.get("is_first_item", True)

        if b_no and b_no not in seen_bills_sub:
            seen_bills_sub.add(b_no)
            sub_sale += float(t_bill or 0)
            sub_paid += float(a_paid or 0)
            sub_due += float(b_due or 0)

        fill2 = EVEN_FILL if r2 % 2 == 0 else ODD_FILL
        item_amt = float(row.get("amount") or 0)
        qty = float(row.get("quantity") or 0)
        rate = float(row.get("rate") or 0)

        vals = [
            cname,
            row.get("customer_phone") or "",
            b_no,
            row.get("bill_date") or "",
            row.get("product_name") or "",
            qty,
            row.get("unit") or "kg",
            rate,
            item_amt,
            float(t_bill) if (is_first_item and t_bill not in [None, ""]) else "",
            float(a_paid) if (is_first_item and a_paid not in [None, ""]) else "",
            float(b_due)  if (is_first_item and b_due  not in [None, ""]) else "",
            row.get("payment_method") or "Cash" if is_first_item else "",
            (row.get("payment_status") or "unpaid").upper() if is_first_item else "",
        ]

        for ci, val in enumerate(vals, 1):
            cell = ws2.cell(row=r2, column=ci, value=val)
            cell.font = BODY_FONT; cell.fill = fill2; cell.border = CELL_BORDER
            if ci in [6, 8, 9, 10, 11, 12]:
                if isinstance(val, (int, float)):
                    cell.number_format = NUM_FMT
                cell.alignment = RIGHT
            elif ci in [1, 5]:
                cell.alignment = LEFT
            else:
                cell.alignment = CENTER

        prev_cust = cname
        r2 += 1

    # Last customer subtotal
    if prev_cust is not None:
        ws2.merge_cells(f"A{r2}:I{r2}")
        sub_lbl = ws2.cell(row=r2, column=1, value=f"  ↳ Total for {prev_cust}")
        sub_lbl.font = Font(name="Segoe UI", size=9, bold=True, color="1E40AF")
        sub_lbl.fill = PatternFill("solid", fgColor="EFF6FF")
        sub_lbl.alignment = LEFT; sub_lbl.border = CELL_BORDER

        for ci, val in enumerate([sub_sale, sub_paid, sub_due], start=10):
            cell = ws2.cell(row=r2, column=ci, value=val)
            cell.font = Font(name="Segoe UI", size=9, bold=True, color="1E40AF")
            cell.fill = PatternFill("solid", fgColor="EFF6FF")
            cell.number_format = NUM_FMT; cell.alignment = RIGHT; cell.border = CELL_BORDER

        for ci in [13, 14]:
            c_end = ws2.cell(row=r2, column=ci, value="")
            c_end.fill = PatternFill("solid", fgColor="EFF6FF"); c_end.border = CELL_BORDER
        r2 += 1

    # Grand Total Row for Sheet 2
    ws2.merge_cells(f"A{r2}:I{r2}")
    lbl2 = ws2.cell(row=r2, column=1, value="GRAND TOTAL")
    lbl2.font = TOTAL_FONT; lbl2.fill = TOTAL_FILL
    lbl2.alignment = RIGHT; lbl2.border = CELL_BORDER

    for ci, val in enumerate([grand_sale, grand_received, grand_pending], start=10):
        cell = ws2.cell(row=r2, column=ci, value=val)
        cell.font = TOTAL_FONT; cell.fill = TOTAL_FILL
        cell.number_format = NUM_FMT; cell.alignment = RIGHT
        cell.border = CELL_BORDER

    for ci in [13, 14]:
        c_end = ws2.cell(row=r2, column=ci, value="")
        c_end.fill = TOTAL_FILL; c_end.border = CELL_BORDER

    ws2.row_dimensions[r2].height = 24

    w2 = [24, 14, 16, 12, 18, 10, 8, 12, 14, 16, 14, 14, 14, 12]
    for ci, w in enumerate(w2, 1):
        ws2.column_dimensions[get_column_letter(ci)].width = w

    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    return buf


def build_daily_by_product(report_data: Union[Dict[str, Any], List[Dict]], report_date: str) -> io.BytesIO:
    """
    Builds a Daily Sales & Product Report with accurate product-wise sales,
    average rates, quantities, and proportional payment allocation.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "By Product"

    if isinstance(report_data, dict):
        products = report_data.get("products", [])
        totals = report_data.get("totals", {})
    else:
        products = report_data
        totals = {
            "total_sales": sum(float(p.get("total_amount") or 0) for p in products),
            "total_received": sum(float(p.get("amount_paid") or 0) for p in products),
            "total_pending": sum(float(p.get("balance_due") or 0) for p in products),
            "total_qty": sum(float(p.get("total_qty") or 0) for p in products),
        }

    grand_sale = float(totals.get("total_sales") or sum(float(p.get("total_amount") or 0) for p in products))
    grand_received = float(totals.get("total_received") or sum(float(p.get("amount_paid") or 0) for p in products))
    grand_pending = float(totals.get("total_pending") or sum(float(p.get("balance_due") or 0) for p in products))
    grand_qty = float(totals.get("total_qty") or sum(float(p.get("total_qty") or 0) for p in products))
    collection_rate = (grand_received / grand_sale * 100) if grand_sale > 0 else 0.0

    # Banner
    ws.merge_cells("A1:K1")
    c = ws["A1"]
    c.value = f"DAILY SALES & PRODUCT REPORT   |   {report_date}"
    c.font = HEADER_FONT; c.fill = HEADER_FILL; c.alignment = CENTER
    ws.row_dimensions[1].height = 32

    # KPI Cards (Rows 3 & 4)
    kpis = [
        ("TOTAL SALES", f"₹{grand_sale:,.2f}", "16A34A"),
        ("RECEIVED AMOUNT", f"₹{grand_received:,.2f}", "2563EB"),
        ("PENDING BALANCE", f"₹{grand_pending:,.2f}", "DC2626"),
        ("COLLECTION RATE", f"{collection_rate:.1f}%", "7C3AED"),
    ]
    _apply_kpi_cards(ws, kpis, start_row=3, num_cols=8)
    ws.row_dimensions[3].height = 18
    ws.row_dimensions[4].height = 24

    # Column Headers (Row 6)
    headers = [
        "#", "Product Name", "Unit", "Qty Sold", "Purchase Price (₹)", "Avg Sell Rate (₹)",
        "Total Sale (₹)", "Received (₹)", "Pending (₹)", "Profit / Loss (₹)", "Share %"
    ]
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=6, column=ci, value=h)
        cell.font = SUBHDR_FONT; cell.fill = SUBHDR_FILL
        cell.alignment = CENTER; cell.border = CELL_BORDER
    ws.row_dimensions[6].height = 22

    # Data Rows
    current_row = 7
    for idx, p in enumerate(products, 1):
        fill = EVEN_FILL if current_row % 2 == 0 else ODD_FILL
        sale = float(p.get("total_amount") or p.get("total_sale") or 0)
        received = float(p.get("amount_paid") or 0)
        pending = float(p.get("balance_due") or 0)
        qty = float(p.get("total_qty") or p.get("quantity") or 0)
        pp = p.get("purchase_price")
        avg_rate = float(p.get("avg_rate") or (sale / qty if qty > 0 else 0))
        pl = p.get("profit_loss")
        share = (sale / grand_sale * 100) if grand_sale > 0 else 0.0

        ws.cell(row=current_row, column=1, value=idx).alignment = CENTER
        ws.cell(row=current_row, column=2, value=p.get("product_name") or "—").alignment = LEFT
        ws.cell(row=current_row, column=3, value=p.get("unit") or "kg").alignment = CENTER

        qc = ws.cell(row=current_row, column=4, value=qty)
        qc.number_format = NUM_FMT; qc.alignment = RIGHT

        ppc = ws.cell(row=current_row, column=5, value=float(pp) if pp is not None else "—")
        if pp is not None:
            ppc.number_format = NUM_FMT; ppc.alignment = RIGHT
            ppc.font = Font(name="Segoe UI", size=9, color="E11D48")
        else:
            ppc.alignment = CENTER

        rc = ws.cell(row=current_row, column=6, value=avg_rate)
        rc.number_format = NUM_FMT; rc.alignment = RIGHT

        sc = ws.cell(row=current_row, column=7, value=sale)
        sc.number_format = NUM_FMT; sc.alignment = RIGHT
        sc.font = Font(name="Segoe UI", size=9, bold=True, color="166534")

        pc = ws.cell(row=current_row, column=8, value=received)
        pc.number_format = NUM_FMT; pc.alignment = RIGHT
        pc.font = Font(name="Segoe UI", size=9, bold=True, color="1E40AF")

        dc = ws.cell(row=current_row, column=9, value=pending)
        dc.number_format = NUM_FMT; dc.alignment = RIGHT
        dc.font = Font(name="Segoe UI", size=9, bold=True, color="991B1B")

        plc = ws.cell(row=current_row, column=10, value=float(pl) if pl is not None else "—")
        if pl is not None:
            plc.number_format = NUM_FMT; plc.alignment = RIGHT
            plc.font = Font(name="Segoe UI", size=9, bold=True, color="166534" if pl >= 0 else "991B1B")
        else:
            plc.alignment = CENTER

        shc = ws.cell(row=current_row, column=11, value=f"{share:.1f}%")
        shc.alignment = CENTER

        for ci in range(1, 12):
            cell = ws.cell(row=current_row, column=ci)
            cell.border = CELL_BORDER
            cell.fill = fill
            if ci not in [5, 7, 8, 9, 10]:
                cell.font = BODY_FONT

        ws.row_dimensions[current_row].height = 20
        current_row += 1

    # Grand Total Row
    ws.merge_cells(f"A{current_row}:C{current_row}")
    lbl = ws.cell(row=current_row, column=1, value="GRAND TOTAL")
    lbl.font = TOTAL_FONT; lbl.fill = TOTAL_FILL
    lbl.alignment = RIGHT; lbl.border = CELL_BORDER

    # Total Qty
    t_q = ws.cell(row=current_row, column=4, value=grand_qty)
    t_q.font = TOTAL_FONT; t_q.fill = TOTAL_FILL; t_q.number_format = NUM_FMT; t_q.alignment = RIGHT; t_q.border = CELL_BORDER

    # Empty Purchase Price and Avg Rate
    for ci in [5, 6]:
        e_c = ws.cell(row=current_row, column=ci, value="")
        e_c.fill = TOTAL_FILL; e_c.border = CELL_BORDER

    # Totals for Sale, Received, Pending
    for ci, val in enumerate([grand_sale, grand_received, grand_pending], start=7):
        cell = ws.cell(row=current_row, column=ci, value=val)
        cell.font = TOTAL_FONT; cell.fill = TOTAL_FILL
        cell.number_format = NUM_FMT; cell.alignment = RIGHT
        cell.border = CELL_BORDER

    # Total Profit / Loss
    has_p_pl = any(p.get("profit_loss") is not None for p in products)
    grand_p_pl = sum(p["profit_loss"] for p in products if p.get("profit_loss") is not None)
    pl_cell = ws.cell(row=current_row, column=10, value=grand_p_pl if has_p_pl else "—")
    pl_cell.font = TOTAL_FONT; pl_cell.fill = TOTAL_FILL; pl_cell.border = CELL_BORDER
    if has_p_pl:
        pl_cell.number_format = NUM_FMT; pl_cell.alignment = RIGHT
    else:
        pl_cell.alignment = CENTER

    # Total Share (100.0%)
    sh_tot = ws.cell(row=current_row, column=11, value="100.0%")
    sh_tot.font = TOTAL_FONT; sh_tot.fill = TOTAL_FILL; sh_tot.alignment = CENTER; sh_tot.border = CELL_BORDER

    ws.row_dimensions[current_row].height = 24

    widths = [6, 26, 10, 14, 18, 16, 18, 18, 18, 18, 12]
    for ci, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    return buf


def build_ledger_page_excel(draft_data: Dict[str, Any], page_date: str) -> io.BytesIO:
    """
    Generate multi-sheet Excel for a digitized ledger page.

    Expects cells pre-enriched by ledger_service.resolve_ledger_rates():
        resolved_product_code, resolved_product_name, resolved_unit,
        resolved_rate, resolved_amount, rate_source, mapping_issue

    Falls back gracefully to raw Gemini qty/rate for old/unresolved drafts.
    """
    wb = Workbook()

    if not isinstance(draft_data, dict):
        draft_data = {}

    col_codes = draft_data.get("column_codes")
    if not col_codes or not isinstance(col_codes, list):
        col_codes = ["M", "R", "B", "P", "K", "T", "JB"]

    rows = draft_data.get("rows", [])
    if not isinstance(rows, list):
        rows = []

    # Sheet 1: Ledger Matrix
    ws1 = wb.active
    ws1.title = "Ledger Grid"

    max_col_letter = get_column_letter(max(len(col_codes) + 3, 6))
    ws1.merge_cells(f"A1:{max_col_letter}1")
    t1 = ws1["A1"]
    t1.value = f"DIGITIZED LEDGER GRID — {page_date or 'UNDATED'}"
    t1.font = HEADER_FONT; t1.fill = HEADER_FILL; t1.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 28

    # ── Sheet 1 header: Customer | Col codes (qty display) | Total Amount | Notes
    headers = ["#", "Customer Name"] + col_codes + ["Total Amount (₹)", "Notes / Tags"]
    for ci, h in enumerate(headers, 1):
        cell = ws1.cell(row=2, column=ci, value=h)
        cell.font = SUBHDR_FONT; cell.fill = SUBHDR_FILL; cell.alignment = Alignment(horizontal="center"); cell.border = CELL_BORDER
    ws1.row_dimensions[2].height = 22

    needs_review_items = []
    tot_amt_col = len(col_codes) + 3
    notes_col = len(col_codes) + 4

    for ri, row in enumerate(rows, start=3):
        if not isinstance(row, dict):
            continue
        fill = EVEN_FILL if ri % 2 == 0 else ODD_FILL
        c_name = row.get("customer_name_raw") or row.get("customer_name") or ""
        ws1.cell(row=ri, column=1, value=ri-2).fill = fill
        ws1.cell(row=ri, column=2, value=c_name).fill = fill

        cells_dict = row.get("cells", {})
        if not isinstance(cells_dict, dict):
            cells_dict = {}

        row_total_amt = 0.0

        for c_idx, code in enumerate(col_codes, start=3):
            cell_raw = cells_dict.get(code, {})
            if isinstance(cell_raw, (int, float, str)):
                cell_data = {"quantity": cell_raw}
            elif isinstance(cell_raw, dict):
                cell_data = cell_raw
            else:
                cell_data = {}

            qty = cell_data.get("quantity")
            # Use resolved_rate if available (DB-authoritative), else fall back to Gemini rate
            resolved_rate = cell_data.get("resolved_rate")
            resolved_amount = cell_data.get("resolved_amount")
            gemini_rate = cell_data.get("rate")
            display_rate = resolved_rate if resolved_rate is not None else gemini_rate
            tag = cell_data.get("tag")
            conf = cell_data.get("confidence")
            circled = cell_data.get("circled_value")

            if resolved_amount is not None:
                try:
                    row_total_amt += float(resolved_amount)
                except (ValueError, TypeError):
                    pass
            elif qty is not None and display_rate is not None:
                try:
                    row_total_amt += float(qty) * float(display_rate)
                except (ValueError, TypeError):
                    pass

            txt = ""
            if qty is not None and display_rate is not None:
                txt = f"{qty} / {display_rate}"
            elif qty is not None:
                txt = str(qty)
            if tag:
                txt += f" ({tag})"

            c_elem = ws1.cell(row=ri, column=c_idx, value=txt or "-")
            c_elem.fill = fill
            c_elem.alignment = Alignment(horizontal="center")
            c_elem.border = CELL_BORDER

            # Confidence / circled / question-mark review flags
            is_low_conf = False
            try:
                if conf is not None and float(conf) < 0.8:
                    is_low_conf = True
            except (ValueError, TypeError):
                pass

            mapping_issue = cell_data.get("mapping_issue")

            if is_low_conf or circled or (qty is not None and str(qty).strip() == "?"):
                needs_review_items.append({
                    "row": ri - 2,
                    "customer": c_name,
                    "column_code": code,
                    "product_code": cell_data.get("resolved_product_code") or "—",
                    "product_name": cell_data.get("resolved_product_name") or "—",
                    "quantity": qty,
                    "rate": display_rate,
                    "confidence": conf,
                    "circled": circled,
                    "reason": "Low OCR confidence" if is_low_conf else "Circled / Flagged item"
                })

            # Also flag mapping / rate problems
            if mapping_issue and qty is not None and str(qty).strip() not in ("", "0"):
                needs_review_items.append({
                    "row": ri - 2,
                    "customer": c_name,
                    "column_code": code,
                    "product_code": cell_data.get("resolved_product_code") or "—",
                    "product_name": cell_data.get("resolved_product_name") or "—",
                    "quantity": qty,
                    "rate": "—",
                    "confidence": conf,
                    "circled": circled,
                    "reason": mapping_issue
                })

        # Customer Total Amount cell in Sheet 1
        tot_cell = ws1.cell(row=ri, column=tot_amt_col, value=round(row_total_amt, 2) if row_total_amt > 0 else "-")
        tot_cell.fill = fill
        tot_cell.font = TOTAL_FONT
        tot_cell.border = CELL_BORDER
        if row_total_amt > 0:
            tot_cell.number_format = NUM_FMT
            tot_cell.alignment = RIGHT
        else:
            tot_cell.alignment = CENTER

        ws1.cell(row=ri, column=notes_col, value=row.get("notes") or "").fill = fill

    # Sheet 1 Bottom Total Row
    if rows:
        tot_ri = len(rows) + 3
        ws1.row_dimensions[tot_ri].height = 24
        for ci in range(1, notes_col + 1):
            c = ws1.cell(row=tot_ri, column=ci)
            c.fill = TOTAL_FILL
            c.font = TOTAL_FONT
            c.border = CELL_BORDER

        ws1.cell(row=tot_ri, column=2, value="GRAND TOTAL").alignment = Alignment(horizontal="right", vertical="center")
        tot_amt_cell = ws1.cell(row=tot_ri, column=tot_amt_col, value=f"=SUM({get_column_letter(tot_amt_col)}3:{get_column_letter(tot_amt_col)}{tot_ri-1})")
        tot_amt_cell.number_format = NUM_FMT
        tot_amt_cell.alignment = RIGHT
        tot_amt_cell.font = TOTAL_FONT

    # ── Sheet 2: Flat Entries List ─────────────────────────────────────────
    ws2 = wb.create_sheet(title="Entries Flat List")
    # Updated headers: Column Code + Product Code + Product Name + Unit columns added
    headers2 = [
        "Date", "Customer",
        "Column", "Product Code", "Product Name", "Unit",
        "Quantity", "Rate (DB)", "Amount",
        "Tag", "Circled Val", "Confidence", "Rate Source"
    ]
    for ci, h in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=ci, value=h)
        cell.font = SUBHDR_FONT; cell.fill = SUBHDR_FILL; cell.alignment = Alignment(horizontal="center"); cell.border = CELL_BORDER
    ws2.row_dimensions[1].height = 22

    e_row = 2
    for row in rows:
        if not isinstance(row, dict):
            continue
        c_name = row.get("customer_name_raw") or row.get("customer_name") or ""
        cells_dict = row.get("cells", {})
        if not isinstance(cells_dict, dict):
            continue
        for code, cell_raw in cells_dict.items():
            if isinstance(cell_raw, (int, float, str)):
                c_data = {"quantity": cell_raw}
            elif isinstance(cell_raw, dict):
                c_data = cell_raw
            else:
                continue

            qty = c_data.get("quantity")
            if qty is not None and str(qty).strip() != "":
                # Prefer DB-resolved values; fall back gracefully for old drafts
                resolved_rate = c_data.get("resolved_rate")
                resolved_amount = c_data.get("resolved_amount")
                resolved_product_code = c_data.get("resolved_product_code") or code
                resolved_product_name = c_data.get("resolved_product_name") or "—"
                resolved_unit = c_data.get("resolved_unit") or "—"
                rate_source = c_data.get("rate_source") or "legacy"

                # If no DB rate, fall back to Gemini rate for display only
                if resolved_rate is None:
                    gemini_rate = c_data.get("rate")
                    try:
                        r_num = float(gemini_rate) if gemini_rate is not None else 0.0
                    except (ValueError, TypeError):
                        r_num = 0.0
                    try:
                        q_num = float(qty)
                    except (ValueError, TypeError):
                        q_num = 0.0
                    display_rate = r_num or "—"
                    display_amount = round(q_num * r_num, 2) if r_num else "—"
                else:
                    display_rate = resolved_rate
                    display_amount = resolved_amount if resolved_amount is not None else "—"

                vals = [
                    page_date,
                    c_name,
                    code,                    # Ledger column code (e.g. M)
                    resolved_product_code,   # Product code (e.g. ON001)
                    resolved_product_name,   # Product name (e.g. Onion)
                    resolved_unit,           # Unit (e.g. kg)
                    qty,
                    display_rate,
                    display_amount,
                    c_data.get("tag") or "",
                    c_data.get("circled_value") or "",
                    c_data.get("confidence") or "",
                    rate_source,
                ]
                fill = EVEN_FILL if e_row % 2 == 0 else ODD_FILL
                for ci, v in enumerate(vals, 1):
                    cell = ws2.cell(row=e_row, column=ci, value=v)
                    cell.font = BODY_FONT; cell.fill = fill; cell.border = CELL_BORDER
                    if ci in (7, 8, 9):
                        if isinstance(v, (int, float)):
                            cell.number_format = NUM_FMT
                        cell.alignment = RIGHT
                    elif ci in (1, 2, 4, 5):
                        cell.alignment = LEFT
                    else:
                        cell.alignment = CENTER
                e_row += 1

    # Sheet 2 Bottom Total Summary Row
    if e_row > 2:
        ws2.row_dimensions[e_row].height = 24
        for ci in range(1, 14):
            c = ws2.cell(row=e_row, column=ci)
            c.fill = TOTAL_FILL
            c.font = TOTAL_FONT
            c.border = CELL_BORDER

        tot_label = ws2.cell(row=e_row, column=5, value="TOTAL")
        tot_label.font = TOTAL_FONT
        tot_label.alignment = Alignment(horizontal="right", vertical="center")

        # Total Quantity (Column G / 7)
        qty_cell = ws2.cell(row=e_row, column=7, value=f"=SUM(G2:G{e_row-1})")
        qty_cell.font = TOTAL_FONT
        qty_cell.number_format = NUM_FMT
        qty_cell.alignment = RIGHT

        # Total Amount (Column I / 9)
        amt_cell = ws2.cell(row=e_row, column=9, value=f"=SUM(I2:I{e_row-1})")
        amt_cell.font = TOTAL_FONT
        amt_cell.number_format = NUM_FMT
        amt_cell.alignment = RIGHT
        e_row += 1
    ws3 = wb.create_sheet(title="Needs Review")
    headers3 = [
        "Row #", "Customer", "Column Code", "Product Code", "Product Name",
        "Quantity", "Rate", "Confidence", "Flag / Circled", "Issue Description"
    ]
    for ci, h in enumerate(headers3, 1):
        cell = ws3.cell(row=1, column=ci, value=h)
        cell.font = SUBHDR_FONT; cell.fill = SUBHDR_FILL; cell.alignment = Alignment(horizontal="center"); cell.border = CELL_BORDER
    ws3.row_dimensions[1].height = 22

    if needs_review_items:
        for n_idx, item in enumerate(needs_review_items, start=2):
            fill = EVEN_FILL if n_idx % 2 == 0 else ODD_FILL
            vals = [
                item.get("row", "-"),
                item.get("customer", "-"),
                item.get("column_code", item.get("product", "-")),  # backward compat
                item.get("product_code") or "—",
                item.get("product_name") or "—",
                item["quantity"] if item.get("quantity") is not None else "-",
                item["rate"] if item.get("rate") is not None else "-",
                item["confidence"] if item.get("confidence") is not None else "-",
                str(item["circled"]) if item.get("circled") else "-",
                item.get("reason", "-")
            ]
            for ci, v in enumerate(vals, 1):
                cell = ws3.cell(row=n_idx, column=ci, value=v)
                cell.font = BODY_FONT; cell.fill = fill; cell.border = CELL_BORDER
    else:
        ws3.cell(row=2, column=1, value="No items requiring manual review. All entries parsed cleanly.").font = BODY_FONT

    # Adjust column widths across sheets
    for ws in [ws1, ws2, ws3]:
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    return buf
